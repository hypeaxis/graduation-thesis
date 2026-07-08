# -*- coding: utf-8 -*-
"""Dựng slide bảo vệ HUST từ template chuẩn (4x3 blue)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

# ROOT = thư mục gốc repo, suy ra từ vị trí script (slide đồ án/build_deck.py) → portable đa máy.
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = os.path.join(ROOT, "slide đồ án/HUST_PPT_template_2022_blue_4x3.pptx")
FIG = os.path.join(ROOT, "Noi_dung_do_an/SOICT_DATN_Research_VIE_Template/figures")
OUT = os.path.join(ROOT, "slide đồ án/Slide_Bao_Ve_HUST.pptx")

# ---- palette ----
NAVY   = RGBColor(0x0B, 0x2A, 0x5E)
RED    = RGBColor(0xC1, 0x20, 0x2E)
GREEN  = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
BLUE   = RGBColor(0x2E, 0x5A, 0xAC)
GRAY   = RGBColor(0x5B, 0x6B, 0x7B)   # GĐ1
S2COL  = BLUE                          # GĐ2
S3COL  = ORANGE                        # GĐ3
TEXT   = RGBColor(0x22, 0x22, 0x22)
MUTED  = RGBColor(0x60, 0x60, 0x60)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT  = RGBColor(0xF2, 0xF4, 0xF8)
LIGHTB = RGBColor(0xDD, 0xE4, 0xF0)
FONT   = "Calibri"

CONTENT_LAYOUT = 8   # 2_Blank : navy title bar + red bar + HUST logo footer
TITLE_LAYOUT   = 1   # 1_Title Slide : navy + red halftone

prs = Presentation(TEMPLATE)

# ---- remove all demo slides ----
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    rId = sldId.get(qn('r:id'))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)

SW, SH = prs.slide_width, prs.slide_height  # EMU

# ================= helpers =================

def content_slide(title):
    s = prs.slides.add_slide(prs.slide_layouts[CONTENT_LAYOUT])
    # strip unused body placeholders (idx 1,2)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx in (1, 2):
            ph._element.getparent().remove(ph._element)
    # title
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text = title
            p = ph.text_frame.paragraphs[0]
            p.font.size = Pt(21); p.font.bold = True
            p.font.color.rgb = WHITE; p.font.name = FONT
            break
    return s

def box(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    return tf

def para(tf, text, size=15, bold=False, color=TEXT, bullet=False, level=0,
         before=2, after=3, align=PP_ALIGN.LEFT, first=False, italic=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.level = level; p.alignment = align
    p.space_before = Pt(before); p.space_after = Pt(after)
    # runs (support inline **bold** segments)
    segs = text.split("**")
    for i, seg in enumerate(segs):
        if seg == "": continue
        r = p.add_run(); r.text = seg
        r.font.size = Pt(size); r.font.name = FONT
        r.font.bold = bold or (i % 2 == 1)
        r.font.italic = italic
        r.font.color.rgb = color
    if bullet:
        _set_bullet(p)
    else:
        _no_bullet(p)
    return p

def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum'):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    if pPr.find(qn('a:buNone')) is None:
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))

def _set_bullet(p, char="•", color="1F4E79"):
    pPr = p._p.get_or_add_pPr()
    pPr.set('indent', '-137160'); pPr.set('marL', '137160')
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    buChar = pPr.makeelement(qn('a:buChar'), {'char': char})
    pPr.append(buFont); pPr.append(buChar)

def rrect(slide, l, t, w, h, fill, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    try:
        sp.adjustments[0] = radius
    except Exception:
        pass
    return sp

def rect_text(slide, l, t, w, h, lines, fill, line=None, anchor=MSO_ANCHOR.MIDDLE,
              align=PP_ALIGN.CENTER, radius=0.08):
    """lines = list of (text, size, bold, color)."""
    sp = rrect(slide, l, t, w, h, fill, line, radius)
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6); tf.margin_right = Pt(6)
    tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
    for i, (txt, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_before = Pt(0); p.space_after = Pt(1)
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT
        _no_bullet(p)
    return sp

def arrow(slide, l, t, w, h, fill=NAVY, shape=MSO_SHAPE.RIGHT_ARROW):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

def add_pic(slide, name, l, t, maxw, maxh, border=True):
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    r = min(maxw / iw, maxh / ih)
    w, h = iw * r, ih * r
    left = l + (maxw - w) / 2
    top = t + (maxh - h) / 2
    pic = slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = LIGHTB; pic.line.width = Pt(0.75)
    return pic

def caption(slide, l, t, w, text):
    tf = box(slide, l, t, w, 0.3)
    para(tf, text, size=9, color=MUTED, italic=True, align=PP_ALIGN.CENTER, first=True)

def table(slide, l, t, w, h, data, col_w=None, header=True, fs=12, head_fs=12,
          head_fill=NAVY, align=None):
    rows, cols = len(data), len(data[0])
    gf = slide.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(h))
    tb = gf.table
    tb.first_row = header; tb.horz_banding = True
    if col_w:
        for j, cw in enumerate(col_w):
            tb.columns[j].width = Inches(cw)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            c = tb.cell(i, j)
            c.margin_left = Pt(5); c.margin_right = Pt(5)
            c.margin_top = Pt(2); c.margin_bottom = Pt(2)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.text = str(val)
            p = c.text_frame.paragraphs[0]
            p.alignment = (align[j] if align else (PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER))
            rr = p.runs[0] if p.runs else p.add_run()
            rr.font.name = FONT
            if i == 0 and header:
                rr.font.size = Pt(head_fs); rr.font.bold = True; rr.font.color.rgb = WHITE
                c.fill.solid(); c.fill.fore_color.rgb = head_fill
            else:
                rr.font.size = Pt(fs); rr.font.color.rgb = TEXT
                c.fill.solid()
                c.fill.fore_color.rgb = WHITE if (i % 2 == 1) else LIGHT
    return tb

def cell_style(tb, i, j, fill=None, color=None, bold=None):
    c = tb.cell(i, j)
    if fill is not None:
        c.fill.solid(); c.fill.fore_color.rgb = fill
    p = c.text_frame.paragraphs[0]
    rr = p.runs[0] if p.runs else p.add_run()
    if color is not None: rr.font.color.rgb = color
    if bold is not None: rr.font.bold = bold

def stage_tag(slide, l, t, text, color):
    """small rounded tag for GĐ label."""
    return rect_text(slide, l, t, 1.5, 0.34, [(text, 11, True, WHITE)], color, radius=0.3)

# ================= SLIDE 1 : Title =================
s = prs.slides.add_slide(prs.slide_layouts[TITLE_LAYOUT])
# remove default title ph to control fully
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
tf = box(s, 1.0, 0.95, 8.0, 0.5)
para(tf, "TRƯỜNG ĐẠI HỌC BÁCH KHOA HÀ NỘI", 14, True, WHITE, align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.0, 1.45, 8.0, 0.4)
para(tf, "ĐỒ ÁN TỐT NGHIỆP", 15, True, RGBColor(0xE0,0x8A,0x90), align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.1, 2.75, 7.8, 1.8)
para(tf, "HỆ THỐNG PHÁT HIỆN XÂM NHẬP MẠNG LAI GHÉP", 26, True, WHITE, align=PP_ALIGN.CENTER, first=True, after=2)
para(tf, "kết hợp Snort và FT-Transformer", 20, False, RGBColor(0xC9,0xD4,0xE6), align=PP_ALIGN.CENTER)
tf = box(s, 1.0, 5.15, 8.0, 1.6)
para(tf, "Sinh viên thực hiện:  [Họ và tên] — MSSV [……]", 14, False, WHITE, align=PP_ALIGN.CENTER, first=True, after=3)
para(tf, "Lớp / Viện:  [……]", 14, False, WHITE, align=PP_ALIGN.CENTER, after=3)
para(tf, "Giảng viên hướng dẫn:  [Học hàm, học vị — Họ tên]", 14, False, WHITE, align=PP_ALIGN.CENTER, after=3)
para(tf, "Hà Nội, [tháng] / [năm]", 12, False, RGBColor(0xC9,0xD4,0xE6), align=PP_ALIGN.CENTER)

# ================= SLIDE 2 : Nội dung =================
s = content_slide("Nội dung trình bày")
items = [
    ("1", "Mục tiêu & bài toán"),
    ("2", "Ba thách thức cốt lõi"),
    ("3", "Phương pháp — ba giai đoạn nghiên cứu"),
    ("4", "Kết quả thực nghiệm"),
    ("5", "Kết luận & hướng phát triển"),
]
y = 1.75
for num, txt in items:
    rect_text(s, 0.9, y, 0.55, 0.55, [(num, 20, True, WHITE)], NAVY, radius=0.5)
    tf = box(s, 1.65, y+0.02, 7.4, 0.5)
    para(tf, txt, 18, True, TEXT, align=PP_ALIGN.LEFT, first=True)
    y += 0.92

# ================= SLIDE 3 : Đặt vấn đề & Mục tiêu =================
s = content_slide("Đặt vấn đề & Mục tiêu")
# left: problem
rect_text(s, 0.5, 1.55, 4.35, 0.45, [("VẤN ĐỀ", 14, True, WHITE)], RED, radius=0.12)
rect_text(s, 0.5, 2.15, 2.1, 1.15,
          [("4,45 triệu USD", 20, True, RED), ("chi phí TB / vụ vi phạm", 11, False, MUTED)],
          LIGHT, line=LIGHTB)
rect_text(s, 2.75, 2.15, 2.1, 1.15,
          [("204 ngày", 20, True, RED), ("thời gian phát hiện TB", 11, False, MUTED)],
          LIGHT, line=LIGHTB)
tf = box(s, 0.5, 3.5, 4.35, 1.9)
para(tf, "Nguồn: Báo cáo IBM 2023", 10, False, MUTED, italic=True, first=True, after=8)
para(tf, "Bài toán NIDS:", 14, True, NAVY, after=2)
para(tf, "phân loại mỗi flow → Benign / loại tấn công", 13, False, TEXT, bullet=True)
para(tf, "yêu cầu xử lý thời gian thực", 13, False, TEXT, bullet=True)
# right: objectives
rect_text(s, 5.15, 1.55, 4.35, 0.45, [("MỤC TIÊU", 14, True, WHITE)], NAVY, radius=0.12)
objs = [
    ("1", "Macro F1 > 0,9 trên CIC-IDS-2017"),
    ("2", "Macro F1 > 0,9 trên Testbed thực"),
    ("3", "Snort + ML phát hiện bổ sung (End-to-End)"),
]
y = 2.2
for num, txt in objs:
    rect_text(s, 5.15, y, 0.5, 0.85, [(num, 18, True, WHITE)], BLUE, radius=0.15)
    rect_text(s, 5.75, y, 3.75, 0.85, [(txt, 13, True, TEXT)], LIGHT, line=LIGHTB,
              align=PP_ALIGN.LEFT)
    y += 1.02

# ================= SLIDE 4 : Ba thách thức =================
s = content_slide("Bài toán & ba thách thức cốt lõi")
cards = [
    ("1. Mất cân bằng cực đoan", "65.000 : 1", "Benign 83–99% · Infiltration <0,01%", RED, GRAY),
    ("2. Covariate shift", "99,55% → 21,93%", "CIC → Testbed WSL2 khi triển khai", RED, S3COL),
    ("3. Giới hạn một tầng", "Snort ⟷ ML", "Snort bỏ sót zero-day · ML bỏ sót 'giống Benign'", NAVY, S2COL),
]
y = 1.55
for title, big, sub, bigcol, barcol in cards:
    rrect(s, 0.5, y, 0.14, 1.15, barcol, radius=0.3)
    rect_text(s, 0.72, y, 8.78, 1.15,
              [(title, 15, True, NAVY), (big, 22, True, bigcol), (sub, 12, False, MUTED)],
              LIGHT, line=LIGHTB, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)
    y += 1.32
tf = box(s, 0.5, y+0.02, 9.0, 0.4)
para(tf, "→ Mỗi thách thức được giải quyết bởi đúng một giai đoạn nghiên cứu.",
     14, True, NAVY, align=PP_ALIGN.CENTER, first=True)

# ================= SLIDE 5 : Lý do chọn FTT =================
s = content_slide("Giới hạn hướng tiếp cận & lý do chọn FT-Transformer")
data = [
    ["Hướng tiếp cận", "Ưu điểm", "Hạn chế"],
    ["Signature (Snort)", "Nhanh, chính xác với tấn công đã biết", "Bỏ sót zero-day / không có dấu hiệu"],
    ["ML thống kê (RF/SVM)", "Không cần viết luật thủ công", "Kém với lớp < 0,01%"],
    ["DL — FT-Transformer", "Attention học tương tác giữa đặc trưng", "Hộp đen, cần đủ dữ liệu"],
]
tb = table(s, 0.5, 1.6, 9.0, 2.3, data, col_w=[2.5, 3.4, 3.1], fs=12.5, head_fs=12.5,
           align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
for j in range(3):
    cell_style(tb, 3, j, fill=LIGHTB, color=NAVY, bold=True)
rect_text(s, 0.5, 4.15, 9.0, 1.35,
          [("Vì sao FT-Transformer?", 14, True, WHITE),
           ("Tấn công như Infiltration chỉ lộ qua TỔ HỢP đặc trưng (đích lạ + upload cao + khuya đêm)",
            13, False, WHITE),
           ("→ Attention phù hợp dữ liệu dạng bảng hơn MLP / Random Forest", 13, True, RGBColor(0xFF,0xD9,0x66))],
          NAVY, align=PP_ALIGN.CENTER)

# ================= SLIDE 6 : Kiến trúc tổng quan =================
s = content_slide("Kiến trúc tổng quan hệ thống")
tf = box(s, 0.5, 1.5, 9.0, 0.3)
para(tf, "Pipeline xử lý End-to-End", 13, True, NAVY, first=True)
flow = [("Traffic\nmạng", GRAY), ("Snort\n(dấu hiệu)", NAVY), ("CICFlowMeter\n(đặc trưng)", NAVY),
        ("FT-Transformer\n(ML đa lớp)", BLUE), ("Alert\nAggregator", RED)]
n = len(flow); x = 0.5; bw = 1.5; gap = 0.38; y = 1.9; bh = 0.95
for i, (txt, col) in enumerate(flow):
    lines = [(ln, 11, True, WHITE) for ln in txt.split("\n")]
    rect_text(s, x, y, bw, bh, lines, col)
    if i < n-1:
        arrow(s, x+bw+0.02, y+bh/2-0.13, gap-0.06, 0.26, fill=RGBColor(0x9A,0xA6,0xB8))
    x += bw + gap
# stage row
tf = box(s, 0.5, 3.25, 9.0, 0.3)
para(tf, "Mạch ba giai đoạn nghiên cứu (mũi tên = kết quả GĐ trước lộ vấn đề GĐ sau)", 13, True, NAVY, first=True)
stages = [("NSL-KDD\nnền tảng · F1 0,68", GRAY), ("CIC-IDS-2017\nMacro F1 0,93", BLUE),
          ("Testbed thực\nMacro F1 0,92", ORANGE)]
x = 1.0; bw = 2.5; gap = 0.75; y = 3.7; bh = 1.0
for i, (txt, col) in enumerate(stages):
    lines = []
    for k, ln in enumerate(txt.split("\n")):
        lines.append((ln, 13 if k == 0 else 11, True if k == 0 else False, WHITE))
    rect_text(s, x, y, bw, bh, lines, col)
    if i < len(stages)-1:
        arrow(s, x+bw+0.05, y+bh/2-0.16, gap-0.1, 0.32, fill=NAVY)
    x += bw + gap
tf = box(s, 0.5, 4.95, 9.0, 0.5)
para(tf, "Các slide sau đi sâu từng khối.", 12, False, MUTED, italic=True, align=PP_ALIGN.CENTER, first=True)

# ================= SLIDE 7 : GĐ1 NSL-KDD =================
s = content_slide("Giai đoạn 1: Mô hình nền tảng (NSL-KDD)")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 1", GRAY)
tf = box(s, 0.5, 2.0, 4.7, 3.4)
para(tf, "Kiến trúc hai tầng:", 14, True, NAVY, first=True, after=2)
para(tf, "Autoencoder Anomaly Gate — lọc Benign, không cần nhãn tấn công", 12.5, False, TEXT, bullet=True)
para(tf, "Stacking Ensemble: FTT + LightGBM → Meta-LR", 12.5, False, TEXT, bullet=True)
para(tf, "Xử lý mất cân bằng:", 14, True, NAVY, before=6, after=2)
para(tf, "CB-Focal Loss · Boost-Minority Val · WeightedRandomSampler", 12.5, False, TEXT, bullet=True)
rect_text(s, 0.5, 4.7, 4.7, 0.7,
          [("Macro F1 = 0,6809", 17, True, GRAY)], LIGHT, line=LIGHTB)
add_pic(s, "fig_nslkdd_perclass.png", 5.4, 1.9, 4.1, 2.9)
caption(s, 5.4, 4.85, 4.1, "Kết quả per-class trên KDDTest+")
rect_text(s, 5.4, 5.15, 4.1, 0.35,
          [("R2L Recall thấp = giới hạn cứng dataset (telnet→pop3)", 10.5, True, RED)],
          RGBColor(0xFD,0xEC,0xEC))

# ================= SLIDE 8 : GĐ2 Two-Stage Cascade =================
s = content_slide("Giai đoạn 2: Two-Stage Cascade (CIC-IDS-2017)")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 2", S2COL)
# flow diagram
y = 2.15; bh = 1.0
rect_text(s, 0.5, y, 1.9, bh, [("Traffic\nCIC-IDS-2017", 12, True, WHITE)], GRAY)
arrow(s, 2.45, y+bh/2-0.13, 0.35, 0.26, fill=RGBColor(0x9A,0xA6,0xB8))
rect_text(s, 2.85, y, 2.3, bh,
          [("Gating Network", 13, True, WHITE), ("nhị phân · ngưỡng 0,85", 10.5, False, WHITE)], BLUE)
arrow(s, 5.2, y+bh/2-0.13, 0.35, 0.26, fill=RGBColor(0x9A,0xA6,0xB8))
rect_text(s, 5.6, y, 2.4, bh,
          [("Expert Network", 13, True, WHITE), ("9 lớp (gộp từ 15 nhãn)", 10.5, False, WHITE)], NAVY)
# benign exit
arrow(s, 3.7, y+bh+0.05, 0.5, 0.45, fill=GREEN, shape=MSO_SHAPE.DOWN_ARROW)
rect_text(s, 4.3, y+bh+0.1, 3.0, 0.5,
          [("Benign ~82,7% thoát pipeline", 12, True, WHITE)], GREEN, align=PP_ALIGN.LEFT)
tf = box(s, 0.5, 4.35, 5.6, 1.15)
para(tf, "Vì sao hai tầng:", 13, True, NAVY, first=True, after=2)
para(tf, "tách 'phát hiện bất thường' khỏi 'phân loại chi tiết' → gradient lớp hiếm không bị 83,4% Benign lấn át",
     12.5, False, TEXT, bullet=True)
para(tf, "Expert chỉ xử lý ~17,3% lưu lượng đáng ngờ", 12.5, False, TEXT, bullet=True)
rect_text(s, 6.3, 4.35, 3.2, 1.15,
          [("Ablation Macro F1", 12, True, WHITE),
           ("1-Stage 0,7831  →  Two-Stage 0,8817", 13, True, RGBColor(0xFF,0xD9,0x66))],
          S2COL)

# ================= SLIDE 9 : GĐ2 HNM =================
s = content_slide("Giai đoạn 2: Hard Negative Mining cho lớp hiếm")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 2", S2COL)
tf = box(s, 0.5, 2.0, 4.7, 1.5)
para(tf, "Ý tưởng — lặp 2 vòng:", 14, True, NAVY, first=True, after=2)
para(tf, "thu mẫu bị phân loại sai (hard negatives)", 12.5, False, TEXT, bullet=True)
para(tf, "thêm vào train với trọng số w_hard = 2,0", 12.5, False, TEXT, bullet=True)
para(tf, "huấn luyện lại → tập trung gradient vào biên quyết định", 12.5, False, TEXT, bullet=True)
rect_text(s, 0.5, 3.6, 2.3, 1.0,
          [("Botnet F1", 12, True, NAVY), ("0,48 → 0,65", 15, True, GREEN), ("+36%", 13, True, GREEN)],
          LIGHT, line=LIGHTB)
rect_text(s, 2.9, 3.6, 2.3, 1.0,
          [("Infiltration F1", 12, True, NAVY), ("0,38 → 0,59", 15, True, GREEN), ("+54,5%", 13, True, GREEN)],
          LIGHT, line=LIGHTB)
rect_text(s, 0.5, 4.75, 4.7, 0.62,
          [("Vượt class weight ×10 (Botnet 0,65 vs 0,52)", 12, True, WHITE)], NAVY)
add_pic(s, "fig_hnm_ablation.png", 5.4, 1.95, 4.1, 3.1)
caption(s, 5.4, 5.05, 4.1, "Tiến triển F1 lớp hiếm qua các vòng HNM")

# ================= SLIDE 10 : GĐ2 Asymmetric Voting =================
s = content_slide("Giai đoạn 2: Asymmetric Voting & kết quả")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 2", S2COL)
tf = box(s, 0.5, 2.0, 4.7, 0.7)
para(tf, "Ensemble: FTT + Random Forest + KNN", 13.5, True, NAVY, first=True, after=1)
para(tf, "(ba mô hình sai trên mẫu khác nhau → bổ sung)", 11.5, False, MUTED)
rect_text(s, 0.5, 2.8, 4.7, 0.85,
          [("Infiltration — luật Ưu tiên", 12.5, True, NAVY),
           ("RF hoặc KNN bỏ phiếu → nhận  (tăng Recall)", 11.5, False, TEXT)],
          LIGHT, line=LIGHTB, align=PP_ALIGN.LEFT)
rect_text(s, 0.5, 3.75, 4.7, 0.85,
          [("Botnet — luật Đồng thuận", 12.5, True, NAVY),
           ("cả 3 cùng phiếu mới nhận  (tăng Precision)", 11.5, False, TEXT)],
          LIGHT, line=LIGHTB, align=PP_ALIGN.LEFT)
add_pic(s, "fig_ensemble_compare.png", 5.4, 1.95, 4.1, 2.75)
rect_text(s, 0.5, 4.8, 9.0, 0.62,
          [("Accuracy 99,55%  ·  Macro F1 = 0,9294   →   ĐẠT MỤC TIÊU 1", 15, True, WHITE)],
          GREEN)

# ================= SLIDE 11 : GĐ3 Covariate Shift =================
s = content_slide("Giai đoạn 3: Chẩn đoán Covariate Shift")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 3", S3COL)
# three drop cards
rect_text(s, 0.5, 1.95, 2.6, 1.15,
          [("CIC (baseline)", 11.5, True, NAVY), ("99,55%", 24, True, GREEN), ("MCC 0,9941", 11, False, MUTED)],
          LIGHT, line=LIGHTB)
arrow(s, 3.15, 2.42, 0.35, 0.26, fill=RED)
rect_text(s, 3.55, 1.95, 2.6, 1.15,
          [("→ Testbed WSL2", 11.5, True, NAVY), ("21,93%", 24, True, RED), ("MCC −0,015 (sai ngược)", 11, False, RED)],
          LIGHT, line=LIGHTB)
arrow(s, 6.2, 2.42, 0.35, 0.26, fill=RED)
rect_text(s, 6.6, 1.95, 2.9, 1.15,
          [("Re-fit Scaler", 11.5, True, NAVY), ("6,87%", 24, True, RED), ("phản trực giác: TỆ HƠN", 11, True, RED)],
          LIGHT, line=LIGHTB)
tf = box(s, 0.5, 3.35, 4.6, 2.0)
para(tf, "Nguyên nhân:", 13.5, True, NAVY, first=True, after=2)
para(tf, "switch vật lý Gigabit → card ảo Hyper-V + NAT", 12, False, TEXT, bullet=True)
para(tf, "đổi packet size & timing → IAT, đếm cờ TCP", 12, False, TEXT, bullet=True)
rect_text(s, 0.5, 4.6, 4.6, 0.75,
          [("Kết luận: fine-tuning không đủ →", 12, True, WHITE),
           ("thu dữ liệu thực + retrain hoàn toàn", 12.5, True, RGBColor(0xFF,0xD9,0x66))],
          NAVY, align=PP_ALIGN.LEFT)
add_pic(s, "fig_da_comparison.png", 5.3, 3.4, 4.2, 1.95)

# ================= SLIDE 12 : GĐ3 Testbed & trung thực =================
s = content_slide("Giai đoạn 3: Testbed thực & tính trung thực đánh giá")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 3", S3COL)
# testbed mini diagram
y = 2.05; bh = 0.75
rect_text(s, 0.5, y, 1.35, bh, [("Kali Linux", 11.5, True, WHITE), ("tấn công", 10, False, WHITE)], RED)
arrow(s, 1.9, y+bh/2-0.11, 0.35, 0.22, fill=RGBColor(0x9A,0xA6,0xB8))
rect_text(s, 2.3, y, 1.35, bh, [("LAN / WiFi", 11.5, True, WHITE), ("router", 10, False, WHITE)], GRAY)
arrow(s, 3.7, y+bh/2-0.11, 0.35, 0.22, fill=RGBColor(0x9A,0xA6,0xB8))
rect_text(s, 4.1, y, 1.9, bh, [("Windows 11 + WSL2", 11, True, WHITE), ("victim (Mirrored Net)", 9.5, False, WHITE)], NAVY)
rect_text(s, 6.15, y, 3.35, bh,
          [("Dataset V8.5", 12, True, WHITE), ("111.825 flows · 5 lớp", 12, True, RGBColor(0xFF,0xD9,0x66))], S3COL)
# two honesty boxes
rect_text(s, 0.5, 3.1, 9.0, 1.05,
          [("Trung thực 1 — PortScan không phân tách trong NAT", 13, True, RED),
           ("< 9% F1 qua MỌI thuật toán (vấn đề dữ liệu) → dùng surrogate CIC Friday PortScan", 12, False, TEXT)],
          LIGHT, line=LIGHTB, align=PP_ALIGN.LEFT)
rect_text(s, 0.5, 4.3, 9.0, 1.15,
          [("Trung thực 2 — Val đa dạng miền tin cậy hơn", 13, True, RED),
           ("val đồng nhất miền ~97%  vs  val đa dạng miền 91,7% (BruteForce F1 chênh 14%)", 12, False, TEXT),
           ("→ chênh lệch do CHẤT LƯỢNG ĐÁNH GIÁ, không phải mô hình  (đóng góp 4)", 12, True, NAVY)],
          LIGHT, line=LIGHTB, align=PP_ALIGN.LEFT)

# ================= SLIDE 13 : GĐ3 V8.5 & hệ thống =================
s = content_slide("Giai đoạn 3: Kết quả V8.5 & Hệ thống End-to-End")
stage_tag(s, 0.5, 1.5, "GIAI ĐOẠN 3", S3COL)
add_pic(s, "fig_v85_perclass.png", 0.5, 1.95, 4.5, 3.1)
caption(s, 0.5, 5.05, 4.5, "Kết quả per-class V8.5 (val đa dạng miền)")
rect_text(s, 5.2, 1.95, 4.3, 0.95,
          [("Macro F1 = 91,7%", 18, True, GREEN),
           ("Balanced Acc 91,1% · MCC 0,865  →  ĐẠT MỤC TIÊU 2", 11.5, True, NAVY)],
          LIGHT, line=LIGHTB)
tf = box(s, 5.2, 3.05, 4.3, 1.0)
para(tf, "PortScan F1 100% (surrogate) · DoS 0,93", 12, False, TEXT, bullet=True, first=True)
para(tf, "BruteForce F1 0,86 (Recall 0,90)", 12, False, TEXT, bullet=True)
rect_text(s, 5.2, 4.15, 4.3, 1.25,
          [("Hệ thống Hybrid tích hợp", 12.5, True, WHITE),
           ("Snort 3 + CICFlowMeter + FTT V8.5 (song song)", 11.5, False, WHITE),
           ("Phần cứng phổ thông: i7 Gen11, 16GB RAM, KHÔNG GPU", 11.5, True, RGBColor(0xFF,0xD9,0x66))],
          NAVY, align=PP_ALIGN.LEFT)

# ================= SLIDE 14 : End-to-End bổ sung =================
s = content_slide("Thử nghiệm End-to-End: tính bổ sung Snort ⟷ ML")
data = [
    ["Loại tấn công", "Snort", "FTT V8.5"],
    ["PortScan (nmap)", "Phát hiện", "Phát hiện"],
    ["BruteForce (hydra)", "Phát hiện", "Phát hiện"],
    ["DoS slowhttptest", "Bỏ sót (HTTP hợp lệ)", "CHỈ ML phát hiện"],
    ["Web Attack SQLi", "Phát hiện", "Phát hiện"],
    ["Web Attack XSS", "CHỈ Snort phát hiện", "Bỏ sót (payload ngắn)"],
]
tb = table(s, 0.75, 1.6, 8.5, 3.05, data, col_w=[3.0, 2.75, 2.75], fs=12.5, head_fs=13,
           align=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER])
cell_style(tb, 3, 1, fill=RGBColor(0xFD,0xEC,0xEC), color=RED, bold=True)
cell_style(tb, 3, 2, fill=RGBColor(0xE7,0xF4,0xE9), color=GREEN, bold=True)
cell_style(tb, 5, 1, fill=RGBColor(0xE7,0xF4,0xE9), color=GREEN, bold=True)
cell_style(tb, 5, 2, fill=RGBColor(0xFD,0xEC,0xEC), color=RED, bold=True)
rect_text(s, 0.75, 4.85, 8.5, 0.62,
          [("Không tấn công nào bị bỏ sót hoàn toàn bởi CẢ HAI tầng  →  ĐẠT MỤC TIÊU 3", 14, True, WHITE)],
          GREEN)

# ================= SLIDE 15 : Demo =================
s = content_slide("Demo hệ thống (Giai đoạn 5 — Replay V8.5)")
ph = rrect(s, 0.7, 1.6, 8.6, 3.05, LIGHT, line=LIGHTB)
tfp = ph.text_frame; tfp.vertical_anchor = MSO_ANCHOR.MIDDLE; tfp.word_wrap = True
p = tfp.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "[ Ảnh chụp màn hình dashboard.html đang phát hiện realtime — kịch bản \"dos\" ]"
r.font.size = Pt(14); r.font.color.rgb = MUTED; r.font.italic = True; r.font.name = FONT
_no_bullet(p)
p2 = tfp.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run(); r2.text = "Nguồn: Final/05_Replay_Detection/dashboard.html"
r2.font.size = Pt(10); r2.font.color.rgb = MUTED; r2.font.name = FONT
_no_bullet(p2)
rect_text(s, 0.7, 4.8, 8.6, 0.68,
          [("REPLAY ≠ LIVE: phát lại corpus flow đã thu (CSV) → suy luận realtime qua WebSocket.", 12, True, WHITE),
           ("Bắt gói trực tiếp (live capture) = hướng phát triển, chưa nằm trong kết quả báo cáo.", 11.5, False, RGBColor(0xFF,0xD9,0x66))],
          RED, align=PP_ALIGN.CENTER)

# ================= SLIDE 16 : Đối chiếu mục tiêu & đóng góp =================
s = content_slide("Đối chiếu mục tiêu & Tổng kết đóng góp")
data = [
    ["Mục tiêu", "Kết quả"],
    ["Macro F1 > 0,9 trên CIC-IDS-2017", "0,9294  ✔"],
    ["Macro F1 > 0,9 trên Testbed thực", "0,917  ✔"],
    ["Snort + ML bổ sung (End-to-End)", "Xác nhận  ✔"],
]
tb = table(s, 0.5, 1.6, 5.0, 1.85, data, col_w=[3.4, 1.6], fs=12, head_fs=12.5,
           align=[PP_ALIGN.LEFT, PP_ALIGN.CENTER])
for i in range(1, 4):
    cell_style(tb, i, 1, color=GREEN, bold=True)
tf = box(s, 5.7, 1.6, 3.8, 3.8)
para(tf, "Bốn đóng góp", 14, True, NAVY, first=True, after=3)
for i, txt in enumerate([
    "Two-Stage Cascade + Asymmetric Voting",
    "Hard Negative Mining 2 vòng cho lớp hiếm",
    "Chẩn đoán covariate shift + Testbed thực + bằng chứng inflate metric",
    "Hybrid IDS End-to-End trên phần cứng phổ thông"], 1):
    para(tf, f"{i}. {txt}", 12.5, False, TEXT, after=6)
tf2 = box(s, 0.5, 3.75, 5.0, 1.6)
para(tf2, "Cả ba mục tiêu đều đạt.", 14, True, GREEN, first=True)

# ================= SLIDE 17 : Hạn chế & hướng phát triển =================
s = content_slide("Hạn chế & Hướng phát triển")
rect_text(s, 0.5, 1.55, 4.4, 0.45, [("HẠN CHẾ", 13, True, WHITE)], RED, radius=0.12)
tf = box(s, 0.5, 2.1, 4.4, 3.3)
for txt in ["Testbed chỉ 5/9 lớp (thiếu DDoS, Botnet, Infiltration, Heartbleed)",
            "Flow-based khó với Infiltration / Botnet",
            "PortScan dùng dữ liệu surrogate",
            "FTT là hộp đen — thiếu explainability",
            "Độ trễ ML ~30s với DoS"]:
    para(tf, txt, 12.5, False, TEXT, bullet=True, first=(txt.startswith("Testbed")), after=6)
rect_text(s, 5.1, 1.55, 4.4, 0.45, [("HƯỚNG PHÁT TRIỂN", 13, True, WHITE)], NAVY, radius=0.12)
tf = box(s, 5.1, 2.1, 4.4, 3.3)
for k, txt in enumerate(["Temporal / session-level (LSTM, Temporal Transformer) cho Botnet, Infiltration, Web Attack",
            "Automated feature engineering",
            "Federated Learning multi-site",
            "Hoàn thiện đường bắt gói trực tiếp (live capture) — đang phát triển"]):
    col = RED if "live capture" in txt else TEXT
    bold = "live capture" in txt
    para(tf, txt, 12.5, bold, col, bullet=True, first=(k == 0), after=6)

# ================= SLIDE 18 : Cảm ơn =================
s = prs.slides.add_slide(prs.slide_layouts[TITLE_LAYOUT])
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
tf = box(s, 1.0, 2.7, 8.0, 1.2)
para(tf, "TRÂN TRỌNG CẢM ƠN", 34, True, WHITE, align=PP_ALIGN.CENTER, first=True, after=4)
para(tf, "Hội đồng đã lắng nghe", 20, False, RGBColor(0xC9,0xD4,0xE6), align=PP_ALIGN.CENTER)
tf = box(s, 1.0, 4.7, 8.0, 0.6)
para(tf, "[Họ tên] — [email/liên hệ]", 14, False, WHITE, align=PP_ALIGN.CENTER, first=True)

# ================= APPENDIX =================
# P1 - divider
s = content_slide("Phụ lục")
tf = box(s, 0.5, 2.8, 9.0, 1.0)
para(tf, "PHỤ LỤC", 30, True, NAVY, align=PP_ALIGN.CENTER, first=True, after=4)
para(tf, "Slide dự phòng cho phần hỏi–đáp", 15, False, MUTED, align=PP_ALIGN.CENTER)

# P2 - PortScan inseparable
s = content_slide("Phụ lục — PortScan không phân tách trong WSL2/NAT")
add_pic(s, "fig_portscan_inseparable.png", 0.5, 1.6, 5.4, 3.7)
tf = box(s, 6.1, 1.9, 3.4, 3.2)
para(tf, "F1 < 9% qua mọi thuật toán:", 13, True, NAVY, first=True, after=3)
for txt in ["FTT (CE / Focal / weight×5): 6,8–8,1%", "Random Forest: 5,2%", "KNN: 4,9%"]:
    para(tf, txt, 12, False, TEXT, bullet=True, after=4)
para(tf, "→ vấn đề DỮ LIỆU, không phải mô hình", 12.5, True, RED, before=6, after=6)
para(tf, "Surrogate CIC Friday PortScan → F1 = 100%", 12.5, True, GREEN, bullet=True)

# P3 - full CIC per-class
s = content_slide("Phụ lục — Kết quả per-class CIC-IDS-2017 (9 lớp)")
data = [["Lớp", "Precision", "Recall", "F1", "Support"],
        ["Benign", "0,9992", "0,9953", "0,9972", "2.031.715"],
        ["DoS", "0,9683", "0,9963", "0,9821", "225.024"],
        ["DDoS", "0,9984", "0,9982", "0,9983", "114.453"],
        ["PortScan", "0,9936", "0,9990", "0,9963", "142.079"],
        ["BruteForce", "0,9473", "0,9989", "0,9724", "12.369"],
        ["Web Attack", "0,9084", "0,9810", "0,9433", "1.950"],
        ["Botnet", "0,7947", "0,6826", "0,7344", "1.758"],
        ["Infiltration", "0,9524", "0,6061", "0,7407", "33"],
        ["Heartbleed", "1,0000", "1,0000", "1,0000", "10"]]
table(s, 0.9, 1.55, 8.2, 4.0, data, col_w=[2.0, 1.55, 1.55, 1.55, 1.55], fs=11.5, head_fs=12)
tf = box(s, 0.9, 5.55, 8.2, 0.35)
para(tf, "Macro F1 = 0,9294 · Accuracy = 99,55%", 12.5, True, NAVY, align=PP_ALIGN.CENTER, first=True)

prs.save(OUT)
print("SAVED:", OUT)
print("N slides:", len(prs.slides._sldIdLst))
